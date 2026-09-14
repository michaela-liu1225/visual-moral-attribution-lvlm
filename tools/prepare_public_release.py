#!/usr/bin/env python3
"""Create the path- and identifier-sanitised CSVs for the public archive.

The script never edits the canonical research outputs. It writes separate
copies under the selected release directory and deliberately excludes raw API
envelopes, participant-level survey records, and image pixels.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
from pathlib import Path


PRIMARY_FILES = (
    "gpt4o_all_results.csv",
    "gpt5.6_all_results.csv",
    "qwen36_flash_all_results.csv",
)

REPEAT_FILES = (
    "gpt4o_subset_repeats.csv",
    "gpt56_subset_repeats.csv",
    "qwen36_subset_repeats.csv",
)


def read_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError(f"Missing CSV header: {path}")
        return list(reader.fieldnames), list(reader)


def write_rows(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def sanitise_primary(source_root: Path, output_root: Path) -> None:
    for filename in PRIMARY_FILES:
        fields, rows = read_rows(source_root / filename)
        fields = [field for field in fields if field != "response_id"]
        for row in rows:
            row.pop("response_id", None)
            if row.get("image_path"):
                row["image_path"] = f"images/{Path(row['image_path']).name}"
        write_rows(output_root / filename, fields, rows)


def sanitise_repeatability(source_root: Path, output_root: Path) -> None:
    for filename in REPEAT_FILES:
        fields, rows = read_rows(source_root / "repeatability_results" / filename)
        for row in rows:
            row["subset_manifest"] = "repeatability_subset_32.csv"
            row["image_path"] = f"images/{row['filename']}"
            row["raw_json_path"] = ""
            row["response_id"] = ""
        write_rows(output_root / "repeatability_results" / filename, fields, rows)


def sanitise_prompt_results(source_root: Path, output_root: Path) -> None:
    source = (
        source_root
        / "prompt_engineering_experiment"
        / "results"
        / "prompt_experiment_results.csv"
    )
    fields, rows = read_rows(source)
    for row in rows:
        row["image_path"] = f"images/{row['filename']}"
        row["raw_json_path"] = ""
        row["response_id"] = ""
        row["execution_id"] = ""
    destination = (
        output_root
        / "prompt_engineering_experiment"
        / "results"
        / "prompt_experiment_results.csv"
    )
    write_rows(destination, fields, rows)


def write_researcher_labels(source_root: Path, output_root: Path) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-dependent
        raise RuntimeError("openpyxl is required to read the source workbook") from exc

    workbook = load_workbook(
        source_root / "merged_model_results_latest.xlsx",
        read_only=True,
        data_only=True,
    )
    worksheet = workbook["Merged"]
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows)
    image_col = header.index("image_id")
    label_col = header.index("Researcher-coded label")
    output_rows = [
        {
            "image_id": str(row[image_col]),
            "researcher_working_label": str(row[label_col]),
        }
        for row in rows
        if row[image_col] is not None
    ]
    write_rows(
        output_root / "data" / "researcher_working_labels.csv",
        ["image_id", "researcher_working_label"],
        output_rows,
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_image_manifest(source_root: Path, output_root: Path) -> None:
    ids = [
        line.strip()
        for line in (source_root / "canonical_127_image_ids.txt").read_text(
            encoding="utf-8"
        ).splitlines()
        if line.strip() and not line.startswith("#")
    ]
    by_id: dict[str, Path] = {}
    for path in (source_root / "images").iterdir():
        if path.is_file() and not path.name.startswith("."):
            if path.stem in by_id:
                raise ValueError(f"Duplicate image identifier: {path.stem}")
            by_id[path.stem] = path
    missing = sorted(set(ids) - set(by_id))
    extra = sorted(set(by_id) - set(ids))
    if missing or extra:
        raise ValueError(f"Image manifest mismatch; missing={missing}, extra={extra}")
    rows = [
        {
            "image_id": image_id,
            "filename": by_id[image_id].name,
            "sha256": sha256_file(by_id[image_id]),
            "size_bytes": str(by_id[image_id].stat().st_size),
        }
        for image_id in ids
    ]
    write_rows(
        output_root / "data" / "image_manifest.csv",
        ["image_id", "filename", "sha256", "size_bytes"],
        rows,
    )


def sanitise_variant_transcript(source_root: Path, output_root: Path) -> None:
    source = source_root / "analysis" / "variant_console_transcript.txt"
    cleaned: list[str] = []
    for line in source.read_text(encoding="utf-8").splitlines():
        if re.search(r"@[^\s%]+(?:\s+[^%]*)?%\s*", line):
            continue
        if re.search(r"^\s*Response ID\s*:", line, flags=re.IGNORECASE):
            continue
        line = re.sub(
            r"\b(?:chatcmpl-|resp_)[A-Za-z0-9_-]+\b", "[redacted]", line
        )
        line = line.replace(str(source_root), "<PROJECT_ROOT>")
        cleaned.append(line)
    destination = output_root / "analysis" / "variant_console_transcript.txt"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text("\n".join(cleaned).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    source_root = args.source_root.resolve()
    output_root = args.output_root.resolve()

    sanitise_primary(source_root, output_root)
    sanitise_repeatability(source_root, output_root)
    sanitise_prompt_results(source_root, output_root)
    write_researcher_labels(source_root, output_root)
    write_image_manifest(source_root, output_root)
    sanitise_variant_transcript(source_root, output_root)


if __name__ == "__main__":
    main()
